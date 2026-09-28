import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Page setup - Margins (2.54 cm = 1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Helper function for headings
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66) # Dark Blue
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
        return p

    def add_paragraph(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.bold = True
        p.add_run(text)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        
        # Set background color to light gray (#F4F4F4)
        shading = parse_xml(r'<w:shd {} w:fill="F4F4F4"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        # Borders: left border thick blue, others none
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            r'<w:tcBorders {}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="005599"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'.format(nsdecls('w'))
        )
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        
        lines = code_text.strip().split('\n')
        for i, line in enumerate(lines):
            if i > 0:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- COVER / TITLE ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("LAPORAN STUDI KASUS OBJECT ORIENTED PROGRAMMING")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("Penerapan Polimorfisme (Method Overriding & Overloading)\npada Sistem Manajemen Operasional Pelabuhan")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(18)
    r_div = p_div.add_run("―" * 40)
    r_div.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
    r_div.font.bold = True

    # --- BAB 1 ---
    add_heading_1("BAB 1: PENDAHULUAN & KONSEP POLIMORFISME")
    
    add_paragraph("Polimorfisme (Polymorphism) berasal dari bahasa Yunani yang berarti 'banyak bentuk'. Dalam Pemrograman Berorientasi Objek (OOP), polimorfisme adalah kemampuan suatu objek atau method untuk memiliki banyak bentuk implementasi yang berbeda meskipun menggunakan nama yang sama.")

    add_heading_2("1.1 Dua Jenis Polimorfisme dalam Java")
    
    add_paragraph("Terjadi ketika beberapa method di dalam kelas yang sama memiliki nama yang sama persis, namun memiliki jumlah, tipe, atau urutan parameter yang berbeda. Keputusan method mana yang dipanggil ditentukan saat kompilasi program.", "1. Polimorfisme Statis (Method Overloading / Compile-Time): ")

    add_paragraph("Terjadi ketika suatu subclass menimpa (mengganti) isi logika dari method yang telah didefinisikan oleh superclass-nya. Method di superclass dan subclass harus memiliki nama dan parameter yang identik. Keputusan eksekusi method ditentukan saat program berjalan berdasarkan bentuk nyata objek yang dibuat.", "2. Polimorfisme Dinamis (Method Overriding / Runtime): ")

    add_heading_2("1.2 Alasan Pemilihan Studi Kasus Pelabuhan")
    add_paragraph("Wilayah pelabuhan merupakan studi kasus yang sangat relevan untuk menggambarkan polimorfisme. Di pelabuhan terdapat berbagai jenis kapal (Kapal Kargo, Kapal Tanker, Kapal Penumpang). Semua kapal tersebut sama-sama merupakan 'Kapal' yang berlabuh dan melakukan bongkar muat, tetapi memiliki mekanisme bongkar muat dan struktur biaya tambat yang berbeda-beda sesuai jenis muatannya.")

    # --- BAB 2 ---
    add_heading_1("BAB 2: ARSITEKTUR KELAS & DESIGN KODE")
    add_paragraph("Sistem Operasional Pelabuhan ini dirancang menggunakan 5 kelas Java utama di dalam paket Polymorphism:")

    # Table of Classes
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header styling
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Nama Class", "Tipe / Peran", "Deskripsi & Penerapan OOP"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        shading = parse_xml(r'<w:shd {} w:fill="003366"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data = [
        ("Kapal", "Superclass (Parent)", "Kelas induk. Menyediakan method yang di-override (hitungBiayaTambat, prosesBongkarMuat) serta variasi Method Overloading (cetakManifest)."),
        ("KapalKargo", "Subclass (Child)", "Turunan Kapal. Meng-override perhitungan biaya tambat (+ sewa derek/crane) dan prosedur bongkar muat kontainer."),
        ("KapalTanker", "Subclass (Child)", "Turunan Kapal. Meng-override biaya tambat (+ hazard safety fee minyak/cair) dan prosedur penyedotan pipa cair."),
        ("KapalPenumpang", "Subclass (Child)", "Turunan Kapal. Meng-override biaya tambat (+ pas pelabuhan per penumpang) dan prosedur penurunan penumpang via garbarata."),
        ("MainRoot", "Main Executable Class", "Titik masuk program. Menguji Dynamic Polymorphism (variabel bertipe Kapal menampung instansi subclass) & Static Polymorphism (Overloading).")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            if row_idx % 2 == 0:
                shading = parse_xml(r'<w:shd {} w:fill="F9F9F9"/>'.format(nsdecls('w')))
                row_cells[col_idx]._tc.get_or_add_tcPr().append(shading)
            p = row_cells[col_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- BAB 3 ---
    add_heading_1("BAB 3: SOURCE CODE LENGKAP & DOKUMENTASI")

    # Read Java files from E:\Mimin Adresteia\pertemuan-4\src\Polymorphism
    base_dir = r"E:\Mimin Adresteia\pertemuan-4\src\Polymorphism"
    files = [
        ("3.1 Superclass: Kapal.java", "Kapal.java"),
        ("3.2 Subclass 1: KapalKargo.java", "KapalKargo.java"),
        ("3.3 Subclass 2: KapalTanker.java", "KapalTanker.java"),
        ("3.4 Subclass 3: KapalPenumpang.java", "KapalPenumpang.java"),
        ("3.5 Main Class: MainRoot.java", "MainRoot.java")
    ]

    for title, fname in files:
        add_heading_2(title)
        fpath = os.path.join(base_dir, fname)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_content = f.read()
            add_code_block(code_content)
        else:
            add_paragraph(f"File {fname} tidak ditemukan.", bold_prefix="[Error] ")

    # --- BAB 4 ---
    add_heading_1("BAB 4: HASIL EKSEKUSI & ANALISIS OUTPUT")
    add_paragraph("Berikut adalah tampilan hasil eksekusi program ketika MainRoot.java dijalankan di terminal:")

    output_text = """==========================================================
    SISTEM MANAJEMEN OPERASIONAL PELABUHAN TANJUNG PRIOK
==========================================================

=== 1. PROSES BONGKAR MUAT (METHOD OVERRIDING) ===
Kapal Kargo [MV Ocean Express] mengoperasikan derek pelabuhan (Gantry Crane) untuk memindahkan 120 kontainer ke area penumpukan (Yard).
Kapal Tanker [MT Pertamina Hero] menghubungkan pipa penyalur keselamatan tinggi untuk menyedot 5000000 liter bahan bakar ke tangki penyimpanan dermaga.
Kapal Penumpang [KMP Dharma Rencana] membuka garbarata/jembatan dermaga untuk menurunkan 450 penumpang dan kendaraan penumpang.

=== 2. HITUNG BIAYA TAMBAT / BERLABUH (OVERRIDING - 3 HARI) ===
Biaya Tambat MV Ocean Express      : Rp 7.500.000
Biaya Tambat MT Pertamina Hero    : Rp 3.500.000
Biaya Tambat KMP Dharma Rencana : Rp 6.000.000

----------------------------------------------------------

=== 3. CETAK MANIFEST KAPAL (METHOD OVERLOADING) ===

[Format 1: Tanpa Parameter]
--- MANIFEST KAPAL ---
Nama Kapal       : MV Ocean Express
No. Registrasi   : KG-8891
Kapasitas (Ton)  : 15000.0 Ton

[Format 2: Dengan Nama Kapten]
--- MANIFEST KAPAL ---
Nama Kapal       : MT Pertamina Hero
No. Registrasi   : TK-4021
Kapasitas (Ton)  : 25000.0 Ton
Kapten Kapal     : Capt. Hendra Wijaya

[Format 3: Dengan Nama Kapten & Jumlah ABK]
--- MANIFEST KAPAL ---
Nama Kapal       : KMP Dharma Rencana
No. Registrasi   : PN-1044
Kapasitas (Ton)  : 5000.0 Ton
Kapten Kapal     : Capt. Agus Santoso
Jumlah ABK       : 35 orang

=========================================================="""

    add_code_block(output_text)

    add_heading_2("4.1 Analisis Hasil")
    add_paragraph("Meskipun ketiga variabel (kargo, tanker, penumpang) disimpulkan sebagai tipe induk 'Kapal', saat method prosesBongkarMuat() dan hitungBiayaTambat() dipanggil, Java mengeksekusi logika milik kelas anaknya masing-masing secara dinamis.", "1. Bukti Method Overriding: ")
    add_paragraph("Pemanggilan method cetakManifest() dengan 0 parameter, 1 parameter (nama kapten), dan 2 parameter (nama kapten + ABK) berhasil memilih method yang tepat secara otomatis sesuai argumen yang dikirimkan.", "2. Bukti Method Overloading: ")

    # --- BAB 5 ---
    add_heading_1("BAB 5: KESIMPULAN")
    add_paragraph("Penerapan Polimorfisme pada studi kasus Operasional Pelabuhan terbukti memberikan fleksibilitas tinggi dalam desain perangkat lunak berorientasi objek:")
    add_paragraph("1. Kode program menjadi lebih modular, rapi, dan mudah dikembangkan (extensible). Jika di masa depan ada tipe kapal baru (misalnya Kapal Perang / Kapal Tunda), kita cukup membuat Subclass baru tanpa merusak kode utama di MainRoot.")
    add_paragraph("2. Method Overriding memungkinkan setiap jenis kapal memiliki perilaku spesifik sesuai muatannya, sementara Method Overloading memberikan fleksibilitas dalam format cetak informasi.")

    # Output path
    out_path = os.path.join(base_dir, "Studi_Kasus_Polymorphism_Pelabuhan.docx")
    doc.save(out_path)
    print("SUCCESS:", out_path)

if __name__ == "__main__":
    create_document()
